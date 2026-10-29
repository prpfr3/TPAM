from pathlib import Path
from django.core.exceptions import ObjectDoesNotExist
from django.core.management.base import BaseCommand
from openpyxl import load_workbook

from locations.models import Location, RouteLocation, RouteMap


DATAIO_DIR = Path(r"C:\\Users\\paulf\\OneDrive\\Data\\TPAM")


class Command(BaseCommand):
    help = "Import RouteLocation records from an Excel export."

    def add_arguments(self, parser):
        parser.add_argument(
            "routemap_name",
            help="Name of the RouteMap to import.",
        )

    def handle(self, *args, **options):

        routemap_name = options["routemap_name"]
        excel_file = DATAIO_DIR / f"RouteLocations_{routemap_name}.xlsx"

        if not excel_file.exists():
            self.stderr.write(self.style.ERROR(f"Excel file not found: {excel_file}"))
            return

        self.stdout.write(f"Importing Route Locations from {excel_file}")

        routemap, _ = RouteMap.objects.get_or_create(name=routemap_name)

        created = 0
        updated = 0
        errors = 0

        # ---------------------------
        # READ EXCEL FILE
        # ---------------------------
        wb = load_workbook(excel_file)
        ws = wb.active

        # First row = headers
        headers = [cell.value for cell in ws[1]]

        # Convert rows to dictionaries like DictReader
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = dict(zip(headers, row))

            print(row_dict)

            loc_no = str(row_dict["loc_no"]).strip()

            route_location, was_created = RouteLocation.objects.get_or_create(
                routemap=routemap,
                loc_no=loc_no,
            )

            # linear_reference
            value = str(row_dict.get("linear_reference") or "").strip()
            if value:
                try:
                    float(value)
                    route_location.linear_reference = value
                except ValueError:
                    pass

            # note
            note = row_dict.get("note")
            route_location.note = (note or "").strip()

            # wikislug lookup
            wikislug = str(row_dict.get("location_wikislug") or "").strip()

            if wikislug:
                location = self.find_location(wikislug)

                if location:
                    route_location.location_fk = location
                else:
                    errors += 1
                    self.stderr.write(
                        f"{loc_no}: Location '{wikislug}' not found."
                    )

            try:
                route_location.save()
            except Exception as e:
                errors += 1
                self.stderr.write(
                    f"{loc_no}: Error saving RouteLocation: {e}"
                )
                continue

            if was_created:
                created += 1
            else:
                updated += 1

        self.stdout.write(
            self.style.SUCCESS(
                f"Finished: {created} created, {updated} updated, {errors} errors."
            )
        )

    @staticmethod
    def find_location(slug):
        try:
            return Location.objects.get(wikislug=slug)
        except ObjectDoesNotExist:
            try:
                return Location.objects.get(slug=slug)
            except ObjectDoesNotExist:
                return None
